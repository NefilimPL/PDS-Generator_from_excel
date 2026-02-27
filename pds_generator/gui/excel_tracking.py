import datetime as dt
import hashlib
import json
import logging
import math
import os
import time
from copy import copy

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Protection, PatternFill

from ..number_format import analyze_numeric_value, DEFAULT_DECIMALS

logger = logging.getLogger(__name__)

TRACKING_COLUMN = "__PDS_ROW_TRACKING__"
_EXCEL_CELL_LIMIT = 32767
_HASH_PREFIX = "sha256:"
_CACHE_VERSION = 1
_WARNING_FILL_RGB = "FFF8E1"
_LEGACY_WARNING_FILL_RGB = "FF0000"
_ERROR_FILL_RGB = "FFC7CE"
_LEGACY_ERROR_FILL_RGB = "FFEBEE"
_WARNING_FILL = PatternFill(
    fill_type="solid",
    start_color=f"FF{_WARNING_FILL_RGB}",
    end_color=f"FF{_WARNING_FILL_RGB}",
)
_ERROR_FILL = PatternFill(
    fill_type="solid",
    start_color=f"FF{_ERROR_FILL_RGB}",
    end_color=f"FF{_ERROR_FILL_RGB}",
)
_DEFAULT_FILL = PatternFill()


def is_tracking_column(name):
    return str(name) == TRACKING_COLUMN


def _normalize_value(value):
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    analysis = analyze_numeric_value(value, decimals=DEFAULT_DECIMALS)
    if analysis is not None:
        return analysis["formatted_dot"]
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def _row_signature(values):
    normalized = [_normalize_value(v) for v in values]
    payload = json.dumps(normalized, ensure_ascii=True, separators=(",", ":"))
    if len(payload) > _EXCEL_CELL_LIMIT:
        digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        return f"{_HASH_PREFIX}{digest}"
    return payload


def _combined_signature(parts):
    payload = json.dumps(parts, ensure_ascii=True, separators=(",", ":"))
    if len(payload) > _EXCEL_CELL_LIMIT:
        digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        return f"{_HASH_PREFIX}{digest}"
    return payload


def _update_cell_protection(cell, locked):
    current = cell.protection
    if current is None:
        cell.protection = Protection(locked=locked)
        return True
    if current.locked == locked:
        return False
    try:
        cell.protection = current.copy(locked=locked)
    except AttributeError:
        cell.protection = Protection(
            locked=locked, hidden=getattr(current, "hidden", False)
        )
    return True


def _fill_rgb(fill):
    if fill is None or getattr(fill, "fill_type", None) in (None, "none"):
        return None
    for color in (
        getattr(fill, "fgColor", None),
        getattr(fill, "start_color", None),
        getattr(fill, "end_color", None),
    ):
        if color is None:
            continue
        rgb = getattr(color, "rgb", None)
        if not rgb:
            continue
        text = str(rgb).strip().lstrip("#")
        if len(text) == 8:
            text = text[2:]
        if len(text) != 6:
            continue
        try:
            int(text, 16)
        except ValueError:
            continue
        return text.upper()
    return None


def _is_tracking_warning_fill(fill):
    rgb = _fill_rgb(fill)
    if rgb is None:
        return False
    if rgb == _WARNING_FILL_RGB:
        return True
    if rgb == _LEGACY_WARNING_FILL_RGB:
        # legacy red markers were used in older versions for data warnings
        return True
    return False


def _is_tracking_error_fill(fill):
    rgb = _fill_rgb(fill)
    if rgb is None:
        return False
    if rgb in {_ERROR_FILL_RGB, _LEGACY_ERROR_FILL_RGB}:
        return True
    r = int(rgb[0:2], 16)
    g = int(rgb[2:4], 16)
    b = int(rgb[4:6], 16)
    return r >= 170 and g <= 110 and b <= 110


def _is_tracking_marker_fill(fill):
    return _is_tracking_warning_fill(fill) or _is_tracking_error_fill(fill)


def _clone_fill(fill):
    try:
        return copy(fill)
    except Exception:
        return PatternFill()


def _column_reference_fill(ws, col_idx, check_rows):
    start_row = 2
    end_row = max(start_row, min((ws.max_row or start_row), check_rows + 1))
    for row in range(start_row, end_row + 1):
        fill = ws.cell(row=row, column=col_idx).fill
        if not _is_tracking_marker_fill(fill):
            return _clone_fill(fill)
    return _clone_fill(_DEFAULT_FILL)


def update_tracking_column(
    excel_path,
    dataframes,
    total_rows,
    sheet_fields=None,
    image_fields=None,
    is_image_missing=None,
):
    if not excel_path or not dataframes:
        return set()
    try:
        wb = load_workbook(excel_path)
    except Exception:
        logger.exception("Failed to open Excel file for tracking update: %s", excel_path)
        raise

    changed_rows = set()
    workbook_dirty = False
    image_field_map = {}
    if isinstance(image_fields, dict):
        for sheet, cols in image_fields.items():
            if not sheet:
                continue
            col_set = {
                str(col) for col in (cols or set()) if col is not None and str(col) != TRACKING_COLUMN
            }
            if col_set:
                image_field_map[str(sheet)] = col_set

    for sheet_name, df in dataframes.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]

        if ws["A1"].value != TRACKING_COLUMN:
            ws.insert_cols(1)
            ws["A1"].value = TRACKING_COLUMN
            workbook_dirty = True

        col_dim = ws.column_dimensions.get("A")
        if col_dim is None or not col_dim.hidden:
            ws.column_dimensions["A"].hidden = True
            workbook_dirty = True

        if sheet_fields is None:
            columns = [col for col in df.columns if col != TRACKING_COLUMN]
        else:
            requested = sheet_fields.get(sheet_name, set())
            columns = sorted({col for col in requested if col != TRACKING_COLUMN})
        header_map = {}
        header_map_str = {}
        for cell in ws[1]:
            if cell.value is None:
                continue
            header_map[cell.value] = cell.column
            header_map_str[str(cell.value)] = cell.column
        column_indices = {}
        for col in columns:
            idx = header_map.get(col)
            if idx is None:
                idx = header_map_str.get(str(col))
            column_indices[col] = idx
        image_columns = image_field_map.get(sheet_name, set())
        row_count = len(df)
        existing_rows = max(0, (ws.max_row or 1) - 1)
        check_rows = max(row_count, existing_rows)
        if total_rows is not None:
            check_rows = min(check_rows, total_rows)
        column_reference_fills = {
            idx: _column_reference_fill(ws, idx, check_rows)
            for idx in column_indices.values()
            if idx
        }

        empty_old = 0
        nonempty_new = 0
        for row_idx in range(check_rows):
            if row_idx < row_count:
                row_series = df.iloc[row_idx]
                values = []
                for col in columns:
                    if col in df.columns:
                        value = row_series[col]
                    else:
                        value = ""
                    try:
                        if pd.isna(value):
                            value = ""
                    except Exception:
                        pass
                    if col in image_columns:
                        values.append(value)
                        col_idx = column_indices.get(col)
                        if col_idx:
                            cell = ws.cell(row=row_idx + 2, column=col_idx)
                            text = str(value or "").strip()
                            should_mark = False
                            if text and callable(is_image_missing):
                                try:
                                    should_mark = bool(is_image_missing(text))
                                except Exception:
                                    logger.exception(
                                        "Failed to validate image for %s!%s row %s",
                                        sheet_name,
                                        col,
                                        row_idx + 2,
                                    )
                            if should_mark:
                                if not _is_tracking_error_fill(cell.fill):
                                    cell.fill = _clone_fill(_ERROR_FILL)
                                    workbook_dirty = True
                            else:
                                if _is_tracking_marker_fill(cell.fill):
                                    cell.fill = _clone_fill(
                                        column_reference_fills.get(col_idx, _DEFAULT_FILL)
                                    )
                                    workbook_dirty = True
                        continue
                    analysis = analyze_numeric_value(
                        value, decimals=DEFAULT_DECIMALS
                    )
                    if analysis is None:
                        values.append(value)
                        col_idx = column_indices.get(col)
                        if col_idx:
                            cell = ws.cell(row=row_idx + 2, column=col_idx)
                            if _is_tracking_warning_fill(cell.fill):
                                cell.fill = _clone_fill(
                                    column_reference_fills.get(col_idx, _DEFAULT_FILL)
                                )
                                workbook_dirty = True
                        continue
                    values.append(analysis["formatted_dot"])
                    col_idx = column_indices.get(col)
                    if not col_idx:
                        continue
                    cell = ws.cell(row=row_idx + 2, column=col_idx)
                    if cell.data_type != "f" and analysis["changed"]:
                        cell.value = analysis["numeric"]
                        workbook_dirty = True
                    if analysis["should_mark"]:
                        if not _is_tracking_warning_fill(cell.fill):
                            cell.fill = _clone_fill(_WARNING_FILL)
                            workbook_dirty = True
                    else:
                        if _is_tracking_warning_fill(cell.fill):
                            cell.fill = _clone_fill(
                                column_reference_fills.get(col_idx, _DEFAULT_FILL)
                            )
                            workbook_dirty = True
            else:
                values = ["" for _ in columns]
            new_sig = _row_signature(values)
            if new_sig != "":
                nonempty_new += 1
            cell = ws.cell(row=row_idx + 2, column=1)
            old_sig = cell.value
            if old_sig is None or (isinstance(old_sig, float) and math.isnan(old_sig)):
                old_sig = ""
            if old_sig == "":
                empty_old += 1
            if old_sig != new_sig:
                cell.value = new_sig
                workbook_dirty = True
                if row_idx < total_rows:
                    changed_rows.add(row_idx)

        if check_rows > 0 and empty_old == check_rows and nonempty_new > 0:
            logger.warning(
                "Tracking column empty for sheet %s; all rows may appear changed",
                sheet_name,
            )

        max_row = max(ws.max_row or 1, check_rows + 1)
        max_col = max(ws.max_column or 1, len(columns) + 1)

        for row in range(1, max_row + 1):
            cell = ws.cell(row=row, column=1)
            if _update_cell_protection(cell, True):
                workbook_dirty = True

        for row in range(1, max_row + 1):
            for col in range(2, max_col + 1):
                cell = ws.cell(row=row, column=col)
                if _update_cell_protection(cell, False):
                    workbook_dirty = True

        if not ws.protection.sheet:
            ws.protection.sheet = True
            workbook_dirty = True

    if workbook_dirty:
        try:
            wb.save(excel_path)
        except Exception:
            logger.exception("Failed to save Excel tracking updates: %s", excel_path)
            raise
    return changed_rows


def mark_execution_error_rows(excel_path, sheet_name, failed_rows, total_rows=None):
    if not excel_path or not sheet_name:
        return
    try:
        wb = load_workbook(excel_path)
    except Exception:
        logger.exception("Failed to open Excel file for execution error marking: %s", excel_path)
        return
    if sheet_name not in wb.sheetnames:
        return

    ws = wb[sheet_name]
    max_col = ws.max_column or 1
    target_col = None
    for col_idx in range(1, max_col + 1):
        header = ws.cell(row=1, column=col_idx).value
        if header is None:
            continue
        if str(header) == TRACKING_COLUMN:
            continue
        target_col = col_idx
        break
    if target_col is None:
        return

    check_rows = max(0, int(total_rows or 0))
    if check_rows <= 0:
        check_rows = max(0, (ws.max_row or 1) - 1)
    reference_fill = _column_reference_fill(ws, target_col, check_rows)
    failed_set = {int(idx) for idx in (failed_rows or []) if isinstance(idx, int) and idx >= 0}
    workbook_dirty = False

    for row_idx in range(check_rows):
        cell = ws.cell(row=row_idx + 2, column=target_col)
        if row_idx in failed_set:
            if not _is_tracking_error_fill(cell.fill):
                cell.fill = _clone_fill(_ERROR_FILL)
                workbook_dirty = True
        else:
            if _is_tracking_error_fill(cell.fill):
                cell.fill = _clone_fill(reference_fill)
                workbook_dirty = True

    if not workbook_dirty:
        return

    try:
        wb.save(excel_path)
    except Exception:
        logger.exception("Failed to save execution error markers to %s", excel_path)


def _tracking_cache_paths(excel_path):
    base, _ext = os.path.splitext(excel_path)
    sidecar = f"{base}.pds_tracking.json"
    cache_dir = os.path.join(os.path.expanduser("~"), ".pds_generator_cache")
    name = os.path.basename(excel_path)
    stem = os.path.splitext(name)[0]
    digest = hashlib.sha256(excel_path.encode("utf-8")).hexdigest()[:12]
    fallback = os.path.join(cache_dir, f"{stem}.{digest}.pds_tracking.json")
    return [sidecar, fallback]


def _read_tracking_cache_file(path):
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except FileNotFoundError:
        return {}
    except Exception:
        logger.exception("Failed to read tracking cache %s", path)
        return {}
    if isinstance(data, dict):
        return data
    return {}


def _load_tracking_cache(paths):
    if isinstance(paths, (list, tuple)):
        for path in paths:
            data = _read_tracking_cache_file(path)
            if data:
                return data, path
        return {}, None
    return _read_tracking_cache_file(paths), paths


def _write_tracking_cache(path, sheet_names, rows, columns_map=None, column_signatures=None):
    payload = {
        "version": _CACHE_VERSION,
        "sheets": sheet_names,
        "columns": columns_map or {},
        "rows": rows,
        "column_signatures": column_signatures or {},
        "timestamp": time.time(),
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=True, separators=(",", ":"))


def _resolve_sheet_columns(dataframes, sheet_fields=None):
    columns_map = {}
    for sheet, df in dataframes.items():
        if sheet_fields is None:
            columns = [col for col in df.columns if col != TRACKING_COLUMN]
        else:
            requested = sheet_fields.get(sheet, set())
            columns = sorted({col for col in requested if col != TRACKING_COLUMN})
        columns_map[sheet] = columns
    return columns_map


def _compute_cache_signatures(dataframes, total_rows, sheet_fields=None):
    sheet_names = sorted(dataframes.keys())
    columns_map = _resolve_sheet_columns(dataframes, sheet_fields)
    row_signatures = []
    for row_idx in range(total_rows):
        per_sheet = []
        for sheet in sheet_names:
            df = dataframes.get(sheet)
            if df is None:
                per_sheet.append(_row_signature([]))
                continue
            columns = columns_map.get(sheet, [])
            if row_idx < len(df):
                row_series = df.iloc[row_idx]
                values = [
                    row_series[col] if col in df.columns else ""
                    for col in columns
                ]
            else:
                values = ["" for _ in columns]
            per_sheet.append(_row_signature(values))
        row_signatures.append(_combined_signature(per_sheet))

    column_signatures = {}
    for sheet in sheet_names:
        df = dataframes.get(sheet)
        columns = columns_map.get(sheet, [])
        sigs = {}
        for col in columns:
            values = []
            if df is not None:
                for row_idx in range(total_rows):
                    if row_idx < len(df) and col in df.columns:
                        values.append(df.iloc[row_idx][col])
                    else:
                        values.append("")
            sigs[col] = _row_signature(values)
        column_signatures[sheet] = sigs

    return sheet_names, columns_map, row_signatures, column_signatures


def update_tracking_cache(excel_path, dataframes, total_rows, sheet_fields=None):
    if not excel_path or not dataframes:
        return set(), {}
    cache_paths = _tracking_cache_paths(excel_path)
    previous, cache_path = _load_tracking_cache(cache_paths)
    if not cache_path:
        cache_path = cache_paths[0]
    prev_rows = previous.get("rows") if isinstance(previous, dict) else None
    prev_sheets = previous.get("sheets") if isinstance(previous, dict) else None
    prev_columns = previous.get("columns") if isinstance(previous, dict) else None
    prev_column_signatures = (
        previous.get("column_signatures") if isinstance(previous, dict) else None
    )
    if not isinstance(prev_rows, list):
        prev_rows = []
    if not isinstance(prev_sheets, list):
        prev_sheets = []
    if not isinstance(prev_columns, dict):
        prev_columns = None
    if not isinstance(prev_column_signatures, dict):
        prev_column_signatures = None

    sheet_names, columns_map, new_rows, column_signatures = _compute_cache_signatures(
        dataframes, total_rows, sheet_fields
    )
    changed_rows = set()
    if prev_sheets and prev_sheets != sheet_names:
        changed_rows = set(range(total_rows))
    elif prev_columns is not None and prev_columns != columns_map:
        changed_rows = set(range(total_rows))
    else:
        for idx in range(total_rows):
            if idx >= len(prev_rows) or prev_rows[idx] != new_rows[idx]:
                changed_rows.add(idx)

    changed_columns = {}
    if prev_column_signatures:
        for sheet, sigs in column_signatures.items():
            prev_sigs = prev_column_signatures.get(sheet, {})
            for col, sig in sigs.items():
                if prev_sigs.get(col) != sig:
                    changed_columns.setdefault(sheet, []).append(col)

    try:
        _write_tracking_cache(
            cache_path,
            sheet_names,
            new_rows,
            columns_map,
            column_signatures,
        )
    except Exception:
        logger.exception("Failed to write tracking cache %s", cache_path)
        for fallback in cache_paths:
            if fallback == cache_path:
                continue
            try:
                os.makedirs(os.path.dirname(fallback), exist_ok=True)
                _write_tracking_cache(
                    fallback,
                    sheet_names,
                    new_rows,
                    columns_map,
                    column_signatures,
                )
            except Exception:
                logger.exception("Failed to write tracking cache %s", fallback)
            else:
                cache_path = fallback
                break
    if changed_columns:
        preview = []
        for sheet, cols in changed_columns.items():
            for col in cols:
                preview.append(f"{sheet}:{col}")
                if len(preview) >= 8:
                    break
            if len(preview) >= 8:
                break
        logger.info(
            "Tracking cache changed columns (sample): %s",
            ", ".join(preview),
        )
    return changed_rows, changed_columns
