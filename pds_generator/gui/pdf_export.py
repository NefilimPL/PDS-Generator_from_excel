import logging
import os
import time
import threading
import re
import math
import hashlib
from collections import OrderedDict
from concurrent.futures import ProcessPoolExecutor, as_completed
from contextlib import suppress
from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import requests
from PIL import Image, ImageOps, UnidentifiedImageError
from openpyxl import load_workbook


def _patch_md5_usedforsecurity_compat():
    """Compatibility shim for Python builds without hashlib `usedforsecurity` kwarg."""
    def _wrap_strip_usedforsecurity(func):
        def _compat(*args, **kwargs):
            kwargs.pop("usedforsecurity", None)
            return func(*args, **kwargs)

        return _compat

    md5_func = getattr(hashlib, "md5", None)
    if md5_func is not None:
        try:
            md5_func(b"", usedforsecurity=False)
        except TypeError as exc:
            if "usedforsecurity" in str(exc):
                hashlib.md5 = _wrap_strip_usedforsecurity(md5_func)
        except Exception:
            pass

    new_func = getattr(hashlib, "new", None)
    if new_func is not None:
        try:
            new_func("md5", b"", usedforsecurity=False)
        except TypeError as exc:
            if "usedforsecurity" in str(exc):
                hashlib.new = _wrap_strip_usedforsecurity(new_func)
        except Exception:
            pass

    with suppress(Exception):
        import _hashlib

        openssl_md5 = getattr(_hashlib, "openssl_md5", None)
        if openssl_md5 is None:
            return
        try:
            openssl_md5(b"", usedforsecurity=False)
        except TypeError as exc:
            if "usedforsecurity" in str(exc):
                _hashlib.openssl_md5 = _wrap_strip_usedforsecurity(openssl_md5)
        except Exception:
            pass


_patch_md5_usedforsecurity_compat()

from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.lib.utils import ImageReader
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from tkinter import messagebox

from ..image_auto_zoom import build_auto_zoom_image
from ..number_format import round_numeric_value
from .. import image_index as image_index_utils
from ..layout_dependencies import (
    DEFAULT_GRID_SIZE,
    apply_layout_dependencies,
    normalize_dependencies,
)
from ..pdf_settings import (
    DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
    blend_pdf_image_target_size,
    get_pdf_image_compression_profile,
    normalize_pdf_image_compression_percent,
)
from ..text_layout import fit_text_lines, pdf_font_name, wrap_text_lines
from ..value_sources import (
    FILE_DATE_KIND_MODIFIED,
    VALUE_SOURCE_DEFAULT,
    normalize_file_date_kind,
    normalize_value_source,
    resolve_element_value,
)

from .excel_tracking import (
    TRACKING_COLUMN,
    mark_execution_error_rows,
    update_tracking_column,
    update_tracking_cache,
)

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = (
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
    ".ico",
)

REMOTE_IMAGE_CHECK_LIMIT_WARNING = (
    "Pominięto część zdalnych sprawdzeń URL obrazów (limit 30 unikalnych linków na uruchomienie)."
)
MISSING_IMAGE_ISSUE_ROW_RE = re.compile(r"^Wiersz\s+(\d+):")
EXCEL_DATA_START_ROW = 2
PDF_PAGE_COMPRESSION = 1
PDF_IMAGE_CACHE_LIMIT = 256
PDF_IMAGE_CACHE_REVISION = 3


def _excel_row_number(data_row_idx):
    return int(data_row_idx) + EXCEL_DATA_START_ROW


def _reportlab_md5_compat(*args, **kwargs):
    kwargs.pop("usedforsecurity", None)
    return hashlib.md5(*args, **kwargs)


def _reportlab_digester_compat(value):
    with suppress(Exception):
        from reportlab.lib.utils import isBytes

        payload = value if isBytes(value) else str(value).encode("utf-8")
        return _reportlab_md5_compat(payload).hexdigest()
    payload = value if isinstance(value, (bytes, bytearray)) else str(value).encode("utf-8")
    return _reportlab_md5_compat(payload).hexdigest()


def _patch_reportlab_md5():
    with suppress(Exception):
        from reportlab.pdfbase import pdfdoc
        from reportlab.lib import utils as rl_utils
        from reportlab.pdfgen import canvas as rl_canvas

        if getattr(pdfdoc, "md5", None) is not _reportlab_md5_compat:
            pdfdoc.md5 = _reportlab_md5_compat
        if getattr(rl_utils, "md5", None) is not _reportlab_md5_compat:
            rl_utils.md5 = _reportlab_md5_compat
        if getattr(rl_utils, "_digester", None) is not _reportlab_digester_compat:
            rl_utils._digester = _reportlab_digester_compat
        if getattr(rl_canvas, "_digester", None) is not _reportlab_digester_compat:
            rl_canvas._digester = _reportlab_digester_compat


_patch_reportlab_md5()


def _to_float(value):
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if value is None:
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not text:
        return None
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _extract_rgb(color_obj):
    if not color_obj:
        return None
    rgb = getattr(color_obj, "rgb", None)
    if not rgb:
        return None
    rgb = str(rgb).strip().lstrip("#")
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return None
    try:
        int(rgb, 16)
    except ValueError:
        return None
    return rgb.upper()


def _is_warning_marker_color(color_obj):
    rgb = _extract_rgb(color_obj)
    if rgb:
        if rgb == "FFF8E1":
            return True
        r = int(rgb[0:2], 16)
        g = int(rgb[2:4], 16)
        b = int(rgb[4:6], 16)
        return r >= 170 and g <= 110 and b <= 110
    indexed = getattr(color_obj, "indexed", None)
    return indexed in {10}


def _cell_marked_warning(cell):
    fill = getattr(cell, "fill", None)
    if fill and getattr(fill, "patternType", None) not in (None, "none"):
        if _is_warning_marker_color(getattr(fill, "fgColor", None)) or _is_warning_marker_color(
            getattr(fill, "start_color", None)
        ):
            return True
    font = getattr(cell, "font", None)
    if font and _is_warning_marker_color(getattr(font, "color", None)):
        return True
    return False


def _collect_red_nonpositive_issues(
    app, row_indices, sheet_fields_all, image_sheet_fields=None
):
    if not row_indices or not sheet_fields_all:
        return []

    task_rows = sorted({idx for idx in row_indices if isinstance(idx, int) and idx >= 0})
    if not task_rows:
        return []

    issues = set()
    try:
        wb = load_workbook(app.excel_path, data_only=False, read_only=True)
    except Exception:
        logger.exception("Failed to open workbook for warning-cell validation")
        return ["Nie udało się sprawdzić pól oznaczonych jako ostrzeżenie w Excelu."]

    for sheet_name, columns in sheet_fields_all.items():
        if sheet_name not in wb.sheetnames:
            continue
        df = app.dataframes.get(sheet_name)
        if df is None or df.empty:
            continue
        image_columns = set((image_sheet_fields or {}).get(sheet_name, set()))
        ws = wb[sheet_name]
        try:
            header_row = next(ws.iter_rows(min_row=1, max_row=1))
        except StopIteration:
            continue
        header_indices = {}
        for col_idx, header_cell in enumerate(header_row, start=1):
            header_value = header_cell.value
            if header_value is None:
                continue
            header_indices[str(header_value)] = col_idx

        for col_name in columns:
            if col_name in image_columns:
                continue
            col_idx = header_indices.get(col_name)
            if not col_idx:
                continue
            for row_idx in task_rows:
                if row_idx < 0 or row_idx >= len(df):
                    continue
                excel_row = row_idx + 2
                cell = ws.cell(row=excel_row, column=col_idx)
                if not _cell_marked_warning(cell):
                    continue
                value = df.iloc[row_idx].get(col_name)
                number = _to_float(value)
                if number is None or number > 0:
                    continue
                shown = round_numeric_value(value)
                issues.add(
                    f"Wiersz {_excel_row_number(row_idx)}: {sheet_name}:{col_name} ma wartość {shown} (<= 0) i jest oznaczone kolorem ostrzeżenia."
                )
    return sorted(issues)


def _is_http_value(value):
    text = str(value or "").strip().lower()
    return text.startswith("http://") or text.startswith("https://")


def _issue_row_index(issue_text):
    match = MISSING_IMAGE_ISSUE_ROW_RE.match(str(issue_text or ""))
    if not match:
        return None
    row_no = int(match.group(1))
    if row_no < EXCEL_DATA_START_ROW:
        return None
    return row_no - EXCEL_DATA_START_ROW


def _strip_issue_row_prefix(issue_text):
    text = str(issue_text or "").strip()
    match = MISSING_IMAGE_ISSUE_ROW_RE.match(text)
    if not match:
        return text
    return text[match.end() :].strip()


def _check_remote_image_exists(url, timeout=3):
    response = None
    try:
        response = requests.head(url, allow_redirects=True, timeout=timeout)
        status = response.status_code
        if status in {405} or status >= 400:
            with suppress(Exception):
                response.close()
            response = requests.get(url, stream=True, timeout=timeout)
            status = response.status_code
        return 200 <= status < 400
    except requests.RequestException:
        return False
    finally:
        if response is not None:
            with suppress(Exception):
                response.close()


def _collect_missing_image_issues(app, row_indices):
    image_fields = set(getattr(app, "image_fields", set()) or [])
    if not row_indices or not image_fields:
        return []

    data_rows = sorted({idx for idx in row_indices if isinstance(idx, int) and idx >= 0})
    if not data_rows:
        return []

    issues = set()
    remote_cache = {}
    max_remote_checks = 30

    dynamic_fields = []
    for field in image_fields:
        if ":" not in field:
            continue
        sheet, col = field.split(":", 1)
        dynamic_fields.append((field, sheet, col))

    for row_idx in data_rows:
        row_no = _excel_row_number(row_idx)
        for field, sheet, col in dynamic_fields:
            df = app.dataframes.get(sheet)
            if df is None or row_idx >= len(df):
                continue
            value = df.iloc[row_idx].get(col)
            try:
                if pd.isna(value):
                    continue
            except Exception:
                pass
            text = str(value or "").strip()
            if not text:
                continue
            if _is_http_value(text):
                if text not in remote_cache and len(remote_cache) < max_remote_checks:
                    remote_cache[text] = _check_remote_image_exists(text)
                if text in remote_cache and not remote_cache[text]:
                    issues.add(
                        f"Wiersz {row_no}: nie można pobrać obrazu z URL dla pola {field}: {text}"
                    )
                continue
            if not app.find_local_image(text):
                issues.add(
                    f"Wiersz {row_no}: nie znaleziono pliku obrazu dla pola {field}: {text}"
                )

    static_entries = _collect_static_entries(app)
    for field in image_fields:
        if ":" in field:
            continue
        value = static_entries.get(field, "")
        text = str(value or "").strip()
        if not text:
            continue
        if _is_http_value(text):
            if text not in remote_cache and len(remote_cache) < max_remote_checks:
                remote_cache[text] = _check_remote_image_exists(text)
            if text in remote_cache and not remote_cache[text]:
                issues.add(
                    f"Pole statyczne {field}: nie można pobrać obrazu z URL: {text}"
                )
            continue
        if not app.find_local_image(text):
            issues.add(
                f"Pole statyczne {field}: nie znaleziono pliku obrazu: {text}"
            )

    if len(remote_cache) >= max_remote_checks:
        issues.add(REMOTE_IMAGE_CHECK_LIMIT_WARNING)
    return sorted(issues)


def _resolve_image_sheet_fields(image_fields):
    mapping = {}
    for field in image_fields or []:
        if ":" not in str(field):
            continue
        sheet, col = str(field).split(":", 1)
        if not sheet or not col:
            continue
        mapping.setdefault(sheet, set()).add(col)
    return mapping


def _build_image_missing_checker(app, max_remote_checks=30):
    remote_cache = {}
    local_cache = {}

    def is_missing(value):
        text = str(value or "").strip()
        if not text:
            return False
        if _is_http_value(text):
            exists = remote_cache.get(text)
            if exists is None:
                if len(remote_cache) >= max_remote_checks:
                    return False
                exists = _check_remote_image_exists(text)
                remote_cache[text] = exists
            return not exists
        exists = local_cache.get(text)
        if exists is None:
            exists = bool(app.find_local_image(text))
            local_cache[text] = exists
        return not exists

    return is_missing


def _ui_call(app, func, *args, **kwargs):
    try:
        if hasattr(app, "ui_call"):
            app.ui_call(func, *args, **kwargs)
        else:
            func(*args, **kwargs)
    except Exception:
        logger.exception("Failed to update UI")


def _ui_status(app, text):
    if hasattr(app, "set_status"):
        _ui_call(app, app.set_status, text)


def _ui_counts(app, total_rows=None, total_tasks=None, processed=None, skipped=None):
    if hasattr(app, "set_counts"):
        _ui_call(
            app,
            app.set_counts,
            total_rows=total_rows,
            total_tasks=total_tasks,
            processed=processed,
            skipped=skipped,
        )


def _ui_progress_mode(app, indeterminate):
    if hasattr(app, "set_progress_indeterminate"):
        _ui_call(app, app.set_progress_indeterminate, indeterminate)


def _ui_finish(app, status):
    if hasattr(app, "finish_generation_ui"):
        _ui_call(app, app.finish_generation_ui, status)


def _notify_generation_complete(app, report):
    if hasattr(app, "on_generation_complete"):
        _ui_call(app, app.on_generation_complete, report)


def _cache_image_index_for_app(app, roots, index_data):
    if not index_data:
        return
    roots_key = tuple(roots or [])
    lock = getattr(app, "image_index_lock", None)
    if lock is not None:
        with lock:
            if hasattr(app, "image_index_data"):
                app.image_index_data = index_data
            if hasattr(app, "image_index_roots"):
                app.image_index_roots = roots_key
        return
    if hasattr(app, "image_index_data"):
        app.image_index_data = index_data
    if hasattr(app, "image_index_roots"):
        app.image_index_roots = roots_key


def _is_cancelled(app):
    cancel_event = getattr(app, "cancel_event", None)
    return cancel_event is not None and cancel_event.is_set()


def to_reportlab_color(value):
    try:
        return colors.HexColor(value)
    except (ValueError, TypeError):
        return colors.toColor(value)


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", str(name))
    cleaned = re.sub(r"\s+", "_", cleaned).strip("_")
    return cleaned


def _wrap_text_lines(text, font_name, font_size, max_width):
    return wrap_text_lines(text, font_name, font_size, max_width)


def _fit_text_lines(text, font_name, max_font_size, box_width, box_height, pad=2):
    return fit_text_lines(text, font_name, max_font_size, box_width, box_height, pad=pad)


def _first_data_column(df):
    for col in df.columns:
        if col != TRACKING_COLUMN:
            return col
    return None


def _get_filename_value(df, name_column, idx):
    if df is None or df.empty or not name_column:
        return ""
    if idx >= len(df):
        return ""
    row = df.iloc[idx]
    value = row.get(name_column, "")
    if pd.isna(value):
        return ""
    return value


def _pdf_image_compression_profile_for_app(app):
    return get_pdf_image_compression_profile(
        getattr(
            app,
            "pdf_image_compression_percent",
            DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
        )
    )


def _pdf_image_target_pixels(width_points, height_points, dpi):
    width_px = max(1, int(math.ceil(float(width_points) * float(dpi) / 72.0)))
    height_px = max(1, int(math.ceil(float(height_points) * float(dpi) / 72.0)))
    return width_px, height_px


def _blended_pdf_image_target_size(
    original_width,
    original_height,
    width_points,
    height_points,
    dpi,
    compression_percent,
):
    full_target_width, full_target_height = _pdf_image_target_pixels(
        width_points,
        height_points,
        dpi=dpi,
    )
    return blend_pdf_image_target_size(
        original_width,
        original_height,
        full_target_width,
        full_target_height,
        compression_percent,
    )


def _image_has_transparency(image):
    if image.mode in {"RGBA", "LA"}:
        alpha = image.getchannel("A")
        alpha_min, _alpha_max = alpha.getextrema()
        return alpha_min < 255
    if image.mode == "P":
        transparency = image.info.get("transparency")
        return transparency is not None
    return False


def _normalise_pdf_image(image, width_points, height_points, dpi, compression_percent):
    prepared = ImageOps.exif_transpose(image)
    prepared.load()
    resized_width, resized_height = _blended_pdf_image_target_size(
        prepared.width,
        prepared.height,
        width_points,
        height_points,
        dpi,
        compression_percent,
    )
    resized_size = (max(1, resized_width), max(1, resized_height))
    if prepared.size != resized_size:
        prepared = prepared.resize(resized_size, Image.LANCZOS)
    return prepared


def _encode_png_for_pdf(image):
    png_image = image
    if png_image.mode not in {"1", "L", "LA", "P", "RGB", "RGBA"}:
        png_image = png_image.convert("RGBA" if _image_has_transparency(png_image) else "RGB")
    output = BytesIO()
    png_image.save(output, format="PNG", optimize=True, compress_level=9)
    output.seek(0)
    return output


def _encode_jpeg_for_pdf(image, quality):
    jpeg_image = image.convert("RGB")
    output = BytesIO()
    jpeg_image.save(
        output,
        format="JPEG",
        quality=quality,
        optimize=True,
        progressive=True,
        subsampling=2,
    )
    output.seek(0)
    return output


def _compress_image_for_pdf(
    image,
    width_points,
    height_points,
    compression_percent=None,
    dpi=None,
    jpeg_quality=None,
):
    if compression_percent is None or dpi is None or jpeg_quality is None:
        default_profile = get_pdf_image_compression_profile(
            DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT
        )
        if compression_percent is None:
            compression_percent = default_profile["compression_percent"]
        if dpi is None:
            dpi = default_profile["target_dpi"]
        if jpeg_quality is None:
            jpeg_quality = default_profile["jpeg_quality"]
    prepared = _normalise_pdf_image(
        image,
        width_points,
        height_points,
        dpi=dpi,
        compression_percent=compression_percent,
    )
    png_buffer = _encode_png_for_pdf(prepared)
    if _image_has_transparency(prepared):
        return png_buffer, "PNG", prepared.size

    jpeg_buffer = _encode_jpeg_for_pdf(prepared, quality=jpeg_quality)
    if len(png_buffer.getbuffer()) <= len(jpeg_buffer.getbuffer()):
        return png_buffer, "PNG", prepared.size
    return jpeg_buffer, "JPEG", prepared.size


def _encode_prepared_image_for_pdf(image, jpeg_quality):
    png_buffer = _encode_png_for_pdf(image)
    if _image_has_transparency(image):
        return png_buffer, "PNG", image.size, True

    jpeg_buffer = _encode_jpeg_for_pdf(image, quality=jpeg_quality)
    if len(png_buffer.getbuffer()) <= len(jpeg_buffer.getbuffer()):
        return png_buffer, "PNG", image.size, False
    return jpeg_buffer, "JPEG", image.size, False


def _background_rgb_for_element(element):
    if getattr(element, "bg_visible", False):
        try:
            color = to_reportlab_color(getattr(element, "bg_color", "white"))
            return (
                int(round(float(getattr(color, "red", 1.0)) * 255.0)),
                int(round(float(getattr(color, "green", 1.0)) * 255.0)),
                int(round(float(getattr(color, "blue", 1.0)) * 255.0)),
            )
        except Exception:
            logger.exception("Failed to resolve image background color, using white fallback")
    return (255, 255, 255)


def _flatten_image_to_background(image, background_rgb):
    rgba_image = image.convert("RGBA")
    base = Image.new(
        "RGBA",
        rgba_image.size,
        (
            int(background_rgb[0]),
            int(background_rgb[1]),
            int(background_rgb[2]),
            255,
        ),
    )
    base.alpha_composite(rgba_image)
    return base.convert("RGB")


def _pdf_image_cache_key(
    source_kind,
    source_ref,
    width_points,
    height_points,
    compression_percent,
    auto_zoom,
    background_rgb=None,
):
    return (
        int(PDF_IMAGE_CACHE_REVISION),
        source_kind,
        source_ref,
        int(round(float(width_points) * 10)),
        int(round(float(height_points) * 10)),
        int(compression_percent),
        bool(auto_zoom),
        tuple(background_rgb) if background_rgb is not None else None,
    )


def _pdf_image_cache_for_app(app):
    cache = getattr(app, "_prepared_pdf_image_cache", None)
    if cache is None:
        cache = OrderedDict()
        setattr(app, "_prepared_pdf_image_cache", cache)
    return cache


def _cache_prepared_pdf_image(app, key, prepared_image):
    cache = _pdf_image_cache_for_app(app)
    if key in cache:
        cache.move_to_end(key)
    cache[key] = prepared_image
    while len(cache) > PDF_IMAGE_CACHE_LIMIT:
        cache.popitem(last=False)


def _load_image_source_bytes(app, value_str):
    source_value = str(value_str or "").strip()
    if not source_value:
        return None, None, None

    if source_value.lower().startswith("http"):
        response = None
        try:
            response = requests.get(source_value, timeout=10)
            response.raise_for_status()
            return "remote", source_value, response.content
        finally:
            if response is not None:
                with suppress(Exception):
                    response.close()

    local_path = None
    if hasattr(app, "find_local_image"):
        local_path = app.find_local_image(source_value)
    elif os.path.isfile(source_value):
        local_path = source_value
    if not local_path:
        return None, None, None

    with open(local_path, "rb") as handle:
        return "local", os.path.normcase(os.path.abspath(local_path)), handle.read()


def _prepare_pdf_image(
    app,
    value_str,
    width_points,
    height_points,
    auto_zoom=False,
    background_rgb=None,
):
    source_kind, source_ref, source_bytes = _load_image_source_bytes(app, value_str)
    if not source_kind or not source_bytes:
        return None

    compression_profile = _pdf_image_compression_profile_for_app(app)
    cache_key = _pdf_image_cache_key(
        source_kind,
        source_ref,
        width_points,
        height_points,
        compression_profile["compression_percent"],
        auto_zoom,
        background_rgb,
    )
    cache = _pdf_image_cache_for_app(app)
    cached = cache.get(cache_key)
    if cached is not None:
        cache.move_to_end(cache_key)
        return cached

    with Image.open(BytesIO(source_bytes)) as image:
        image.load()
        prepared_source = image.copy()
    if auto_zoom and compression_profile.get("passthrough"):
        zoomed_image = build_auto_zoom_image(
            prepared_source,
            width_points,
            height_points,
            resize_to_target=False,
        )
        if zoomed_image is not None:
            prepared_source = zoomed_image
    if auto_zoom and background_rgb is not None and _image_has_transparency(prepared_source):
        prepared_source = _flatten_image_to_background(prepared_source, background_rgb)

    if compression_profile.get("passthrough"):
        prepared_image = SimpleNamespace(
            image=prepared_source,
            reader=ImageReader(prepared_source),
            format="ORIGINAL",
            pixel_size=prepared_source.size,
            has_transparency=_image_has_transparency(prepared_source),
        )
    else:
        if auto_zoom:
            target_width_px, target_height_px = _blended_pdf_image_target_size(
                prepared_source.width,
                prepared_source.height,
                width_points,
                height_points,
                dpi=compression_profile["target_dpi"],
                compression_percent=compression_profile["compression_percent"],
            )
            zoomed_image = build_auto_zoom_image(
                prepared_source,
                target_width_px,
                target_height_px,
                resize_to_target=True,
            )
            if zoomed_image is not None:
                prepared_source = zoomed_image
                if background_rgb is not None and _image_has_transparency(prepared_source):
                    prepared_source = _flatten_image_to_background(
                        prepared_source,
                        background_rgb,
                    )
                buffer, encoded_format, pixel_size, has_transparency = _encode_prepared_image_for_pdf(
                    prepared_source,
                    jpeg_quality=compression_profile["jpeg_quality"],
                )
            else:
                buffer, encoded_format, pixel_size = _compress_image_for_pdf(
                    prepared_source,
                    width_points,
                    height_points,
                    compression_percent=compression_profile["compression_percent"],
                    dpi=compression_profile["target_dpi"],
                    jpeg_quality=compression_profile["jpeg_quality"],
                )
                has_transparency = _image_has_transparency(prepared_source)
        else:
            buffer, encoded_format, pixel_size = _compress_image_for_pdf(
                prepared_source,
                width_points,
                height_points,
                compression_percent=compression_profile["compression_percent"],
                dpi=compression_profile["target_dpi"],
                jpeg_quality=compression_profile["jpeg_quality"],
            )
            has_transparency = _image_has_transparency(prepared_source)

        prepared_image = SimpleNamespace(
            buffer=buffer,
            reader=ImageReader(buffer),
            format=encoded_format,
            pixel_size=pixel_size,
            has_transparency=has_transparency,
        )
    _cache_prepared_pdf_image(app, cache_key, prepared_image)
    return prepared_image


def draw_pdf_element(app, c, element, value, x, y):
    value_str = value if isinstance(value, str) else str(value)
    box_width = element.width / app.scale
    box_height = element.height / app.scale
    if getattr(element, "is_image", False) and value_str:
        try:
            prepared_image = _prepare_pdf_image(
                app,
                value_str,
                box_width,
                box_height,
                auto_zoom=getattr(element, "image_auto_zoom", False),
                background_rgb=_background_rgb_for_element(element),
            )
        except (OSError, UnidentifiedImageError, requests.RequestException):
            logger.exception("Failed to prepare image %s for PDF", value_str)
        else:
            if prepared_image is not None:
                c.drawImage(
                    prepared_image.reader,
                    x,
                    y,
                    width=box_width,
                    height=box_height,
                    mask="auto" if getattr(prepared_image, "has_transparency", False) else None,
                )
                return
    if element.bg_visible:
        c.setFillColor(to_reportlab_color(element.bg_color))
        c.rect(
            x,
            y,
            box_width,
            box_height,
            fill=1,
            stroke=0,
        )
    c.setFillColor(to_reportlab_color(element.text_color))
    font_name = pdf_font_name(getattr(element, "bold", False))
    base_font_size = element.font_size / app.scale
    max_font_size = getattr(element, "max_font_size", element.font_size) / app.scale
    auto_fit = getattr(element, "auto_font", True)
    pad = 2 if auto_fit else 0
    if auto_fit:
        font_size, lines = _fit_text_lines(
            value_str, font_name, max_font_size, box_width, box_height, pad=pad
        )
    else:
        font_size = base_font_size
        lines = value_str.splitlines() or [value_str]
    c.setFont(font_name, font_size)
    ascent = pdfmetrics.getAscent(font_name, font_size)
    descent = pdfmetrics.getDescent(font_name, font_size)
    line_height = ascent - descent
    total_height = line_height * len(lines)
    bottom = y + (box_height - total_height) / 2
    baseline_last = bottom - descent
    for idx, line in enumerate(lines):
        baseline = baseline_last + line_height * (len(lines) - 1 - idx)
        if element.align == "center":
            c.drawCentredString(x + box_width / 2, baseline, line)
        elif element.align == "right":
            c.drawRightString(x + box_width - pad, baseline, line)
        else:
            c.drawString(x + pad, baseline, line)


_render_context = None
_render_proxy = None


class RenderAppProxy:
    def __init__(
        self,
        scale,
        excel_dir,
        image_dirs=None,
        image_index_path="",
        pdf_image_compression_percent=DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
    ):
        self.scale = scale
        self.excel_dir = excel_dir or ""
        self.image_dirs = list(image_dirs or [])
        self.pdf_image_compression_percent = normalize_pdf_image_compression_percent(
            pdf_image_compression_percent
        )
        self._image_cache = {}
        self._image_index_data = None
        roots = image_index_utils.normalize_roots([self.excel_dir] + self.image_dirs)
        if roots:
            self._image_index_data = image_index_utils.load_index_for_roots(
                roots,
                index_path=image_index_path or None,
            )

    def find_local_image(self, filename):
        if not filename:
            return None
        if not self.excel_dir and not self.image_dirs:
            return None
        name = str(filename).strip()
        if not name:
            return None
        key = name.lower()
        if key in self._image_cache:
            return self._image_cache[key]

        roots = [self.excel_dir] + list(self.image_dirs or [])
        seen = set()
        search_roots = []
        for root in roots:
            if not root:
                continue
            root = os.path.abspath(root)
            if root not in seen:
                seen.add(root)
                search_roots.append(root)

        stem, _ext = os.path.splitext(os.path.basename(name))
        stem_lower = stem.lower()

        path = None
        if os.path.isabs(name):
            if os.path.isfile(name):
                path = name
            elif stem:
                abs_dir = os.path.dirname(name)
                for ext in IMAGE_EXTENSIONS:
                    candidate = os.path.join(abs_dir, stem + ext)
                    if os.path.isfile(candidate):
                        path = candidate
                        break
        else:
            for root in search_roots:
                candidate = os.path.join(root, name)
                if os.path.isfile(candidate):
                    path = candidate
                    break
                if stem:
                    rel_dir = os.path.dirname(name)
                    if rel_dir:
                        for ext in IMAGE_EXTENSIONS:
                            candidate = os.path.join(root, rel_dir, stem + ext)
                            if os.path.isfile(candidate):
                                path = candidate
                                break
                        if path:
                            break
                    for ext in IMAGE_EXTENSIONS:
                        candidate = os.path.join(root, stem + ext)
                        if os.path.isfile(candidate):
                            path = candidate
                            break
                    if path:
                        break
        if path is None and self._image_index_data:
            path = image_index_utils.find_in_index(self._image_index_data, name)
        if path is None and stem and not self._image_index_data:
            target_name = os.path.basename(name).lower()
            for root in search_roots:
                for current_root, _dirs, files in os.walk(root):
                    for f in files:
                        f_lower = f.lower()
                        if f_lower == target_name:
                            path = os.path.join(current_root, f)
                            break
                        file_stem, file_ext = os.path.splitext(f_lower)
                        if file_stem == stem_lower and file_ext in IMAGE_EXTENSIONS:
                            path = os.path.join(current_root, f)
                            break
                    if path:
                        break
                if path:
                    break
        self._image_cache[key] = path
        return path


def _init_render_context(context):
    global _render_context, _render_proxy
    _patch_md5_usedforsecurity_compat()
    _patch_reportlab_md5()
    _render_context = context
    _render_proxy = RenderAppProxy(
        context["scale"],
        context["excel_dir"],
        context.get("image_dirs", []),
        context.get("image_index_path", ""),
        context.get(
            "pdf_image_compression_percent",
            DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
        ),
    )


def _collect_element_specs(app):
    elements = {}
    order = []
    for name, element in sorted(app.elements.items(), key=lambda kv: kv[1].layer):
        elements[name] = {
            "x": element.x,
            "y": element.y,
            "width": element.width,
            "height": element.height,
            "font_size": element.font_size,
            "max_font_size": getattr(element, "max_font_size", element.font_size),
            "bold": element.bold,
            "text_color": element.text_color,
            "bg_color": element.bg_color,
            "bg_visible": element.bg_visible,
            "align": element.align,
            "auto_font": getattr(element, "auto_font", True),
            "layer": element.layer,
            "is_image": getattr(element, "is_image", False),
            "image_auto_zoom": getattr(element, "image_auto_zoom", False),
            "value_source": normalize_value_source(
                getattr(element, "value_source", VALUE_SOURCE_DEFAULT)
            ),
            "file_date_kind": normalize_file_date_kind(
                getattr(element, "file_date_kind", FILE_DATE_KIND_MODIFIED)
            ),
        }
        order.append(name)
    return elements, order


def _collect_group_specs(app):
    groups = []
    for group in app.groups.values():
        groups.append(
            {
                "x": group.x,
                "y": group.y,
                "width": group.width,
                "height": group.height,
                "fields": list(group.fields),
                "field_pos": {k: (pos[0], pos[1]) for k, pos in group.field_pos.items()},
                "field_conf": {k: dict(conf) for k, conf in group.field_conf.items()},
                "conditions": [tuple(cond) for cond in group.conditions],
            }
        )
    return groups


def _collect_static_entries(app):
    return {name: var.get() for name, var in getattr(app, "static_entries", {}).items()}


def _collect_dependencies(app):
    return normalize_dependencies(getattr(app, "dependencies", []))


def render_single_pdf(task):
    if _render_context is None or _render_proxy is None:
        raise RuntimeError("Render context not initialised")
    _patch_md5_usedforsecurity_compat()
    _patch_reportlab_md5()

    context = _render_context
    idx = task["idx"]
    target_path = task["pdf_path"]
    tmp_path = f"{target_path}.{os.getpid()}.tmp"

    scale = context["scale"]
    page_width = context["page_width"]
    page_height = context["page_height"]
    element_specs = context["elements"]
    element_order = context["element_order"]
    groups = context["groups"]
    conditions = context["conditions"]
    dependencies = context.get("dependencies", [])
    static_entries = context["static_entries"]
    image_fields = set(context.get("image_fields", []))
    elements = {name: SimpleNamespace(**spec) for name, spec in element_specs.items()}

    needed = set(elements.keys())
    for group in groups:
        needed.update(group["fields"])
    needed.update(static_entries.keys())

    row_values = task.get("row_values", {})
    values = {}
    for name in needed:
        if ":" in name:
            value = row_values.get(name, "")
        else:
            value = static_entries.get(name, "")
        value = resolve_element_value(
            value,
            elements.get(name),
            context.get("excel_path", ""),
        )
        if pd.isna(value):
            value = ""
        values[name] = value

    group_field_names = {fname for group in groups for fname in group["fields"]}

    hidden = set()
    for src, tgt in conditions:
        if src in group_field_names or tgt in group_field_names:
            continue
        if pd.isna(values.get(src)) or values.get(src) == "":
            hidden.add(tgt)

    apply_layout_dependencies(
        elements,
        hidden_names=hidden,
        dependencies=dependencies,
        grid_step=context.get("grid_size", DEFAULT_GRID_SIZE) * scale,
    )

    c = pdf_canvas.Canvas(
        tmp_path,
        pagesize=(page_width, page_height),
        pageCompression=PDF_PAGE_COMPRESSION,
    )
    success = False
    try:
        for group in groups:
            g_hidden = set()
            for src, tgt in group["conditions"]:
                if src not in group["fields"] or tgt not in group["fields"]:
                    continue
                if pd.isna(values.get(src)) or values.get(src) == "":
                    g_hidden.add(tgt)
            positions = group["field_pos"]
            columns = {}
            for fname in group["fields"]:
                if fname in hidden or fname in g_hidden:
                    continue
                val = values.get(fname, "")
                if val == "":
                    continue
                conf = group["field_conf"].get(fname, {})
                el = elements.get(fname)
                if not conf and not el:
                    continue
                width = conf.get("width", el.width / scale if el else 0)
                height = conf.get("height", el.height / scale if el else 0)
                x0, y0 = positions.get(fname, (0, 0))
                columns.setdefault(x0, []).append((y0, fname, width, height, conf, el, val))

            placed = []
            for x0 in sorted(columns):
                col_items = columns[x0]
                col_items.sort(key=lambda t: t[0])
                cur_y = 0
                for _, fname, width, height, conf, el, val in col_items:
                    y = cur_y
                    while True:
                        overlap = False
                        for px, py, pw, ph in placed:
                            if (
                                x0 < px + pw
                                and x0 + width > px
                                and y < py + ph
                                and y + height > py
                            ):
                                y = py + ph
                                overlap = True
                                break
                        if not overlap:
                            break
                    if y + height > group["height"] / scale:
                        continue
                    font_size = conf.get(
                        "font_size", el.font_size / scale if el else 12
                    )
                    dummy = SimpleNamespace(
                        width=width * scale,
                        height=height * scale,
                        font_size=font_size * scale,
                        max_font_size=conf.get(
                            "max_font_size",
                            (el.max_font_size / scale if el and hasattr(el, "max_font_size") else font_size),
                        )
                        * scale,
                        bold=conf.get("bold", el.bold if el else False),
                        text_color=conf.get(
                            "text_color", el.text_color if el else "black"
                        ),
                        bg_color=conf.get("bg_color", el.bg_color if el else "white"),
                        bg_visible=conf.get(
                            "bg_visible", el.bg_visible if el else True
                        ),
                        align=conf.get("align", el.align if el else "left"),
                        auto_font=conf.get("auto_font", el.auto_font if el else True),
                        is_image=(
                            conf.get("is_image")
                            if conf is not None and "is_image" in conf
                            else (el.is_image if el and hasattr(el, "is_image") else fname in image_fields)
                        ),
                        image_auto_zoom=(
                            conf.get("image_auto_zoom")
                            if conf is not None and "image_auto_zoom" in conf
                            else (
                                el.image_auto_zoom
                                if el and hasattr(el, "image_auto_zoom")
                                else False
                            )
                        ),
                    )
                    x_pdf = group["x"] / scale + x0
                    y_pdf = page_height - (group["y"] / scale + y + height)
                    draw_pdf_element(_render_proxy, c, dummy, val, x_pdf, y_pdf)
                    placed.append((x0, y, width, height))
                    cur_y = y + height

        for name in element_order:
            element = elements.get(name)
            if not element or name in hidden:
                continue
            val = values.get(name, "")
            x = element.x / scale
            y = page_height - (element.y / scale) - (element.height / scale)
            draw_pdf_element(_render_proxy, c, element, val, x, y)

        c.showPage()
        c.save()
        final_path = target_path
        try:
            os.replace(tmp_path, final_path)
        except OSError:
            logger.exception("Failed to replace %s, trying alternative name", final_path)
            base, ext = os.path.splitext(final_path)
            alt_path = f"{base}_{idx + 1}_{int(time.time())}{ext}"
            try:
                os.replace(tmp_path, alt_path)
            except OSError:
                logger.exception("Failed to move temp PDF to %s", alt_path)
                raise
            else:
                final_path = alt_path
        success = True
        return {"idx": idx, "pdf_path": final_path, "name": task.get("name")}
    finally:
        if not success:
            with suppress(OSError):
                os.remove(tmp_path)


def generate_pds(app):
    report = {
        "status": "started",
        "excel_path": getattr(app, "excel_path", ""),
        "log_path": getattr(app, "runtime_log_path", ""),
        "output_dir": "",
        "total_rows": 0,
        "total_tasks": 0,
        "processed_rows": 0,
        "skipped_rows": 0,
        "new_pdfs": [],
        "updated_pdfs": [],
        "updated_pdf_changes": [],
        "updated_pdf_changes_note": "",
        "errors": [],
        "warnings": [],
        "skipped_image_files": [],
        "skipped_image_global_issues": [],
        "critical_exception": False,
    }

    def finish_now(status_label, status_code):
        if hasattr(app, "finish_generation_ui"):
            app.finish_generation_ui(status_label)
        else:
            _ui_finish(app, status_label)
        report["status"] = status_code
        _notify_generation_complete(app, report)

    if not app.excel_path or not app.dataframes:
        report["errors"].append("Brak danych do generowania.")
        messagebox.showerror("Błąd", "Brak danych do generowania")
        finish_now("Brak danych", "no_data")
        return False

    first_sheet_name, first_df = next(iter(app.dataframes.items()))
    total_rows = len(first_df)
    report["total_rows"] = total_rows
    if total_rows == 0:
        messagebox.showinfo("Info", "Brak wierszy w pliku Excel")
        finish_now("Brak wierszy", "no_rows")
        return False

    _ui_counts(app, total_rows=total_rows)

    element_specs, element_order = _collect_element_specs(app)
    group_specs = _collect_group_specs(app)
    conditions = [tuple(cond) for cond in app.conditions]
    dependencies = _collect_dependencies(app)

    dynamic_fields = set()

    def collect_dynamic(name):
        if isinstance(name, str) and ":" in name:
            dynamic_fields.add(name)

    for name in element_specs.keys():
        collect_dynamic(name)
    for group in group_specs:
        for fname in group["fields"]:
            collect_dynamic(fname)
        for src, tgt in group["conditions"]:
            collect_dynamic(src)
            collect_dynamic(tgt)
    for src, tgt in conditions:
        collect_dynamic(src)
        collect_dynamic(tgt)

    sheet_fields_all = {}
    sheet_fields_tracking = {}
    image_sheet_fields = _resolve_image_sheet_fields(
        getattr(app, "image_fields", set()) or []
    )
    image_missing_checker = _build_image_missing_checker(app)
    excluded_fields = getattr(app, "tracking_excluded", set())
    for field in dynamic_fields:
        sheet, col = field.split(":", 1)
        if col == TRACKING_COLUMN:
            continue
        sheet_fields_all.setdefault(sheet, set()).add(col)
        if field in excluded_fields:
            continue
        sheet_fields_tracking.setdefault(sheet, set()).add(col)

    cache_rows = None
    cache_changed_cols = {}
    cache_row_changes = {}
    excel_rows = None
    tracking_mode = "none"

    try:
        cache_rows, cache_changed_cols, cache_row_changes = update_tracking_cache(
            app.excel_path, app.dataframes, total_rows, sheet_fields=sheet_fields_tracking
        )
    except Exception:
        logger.exception("Failed to update tracking cache")
        report["critical_exception"] = True
        report["errors"].append(
            "Nie udało się zaktualizować lokalnego cache śledzenia zmian."
        )
        cache_rows = None
        cache_changed_cols = {}
        cache_row_changes = {}

    try:
        excel_rows = update_tracking_column(
            app.excel_path,
            app.dataframes,
            total_rows,
            sheet_fields=sheet_fields_tracking,
            image_fields=image_sheet_fields,
            is_image_missing=image_missing_checker,
        )
        tracking_mode = "excel"
    except Exception:
        logger.exception("Failed to update tracking column")
        report["critical_exception"] = True
        report["errors"].append(
            "Nie udało się zaktualizować kolumny kontrolnej w Excelu."
        )
        if cache_rows is None:
            messagebox.showwarning(
                "Uwaga",
                "Nie udało się zaktualizować kolumny kontrolnej w Excelu "
                "(sprawdź, czy plik nie jest otwarty). "
                "Pliki PDF zostaną wygenerowane ponownie.",
            )
        else:
            tracking_mode = "cache"
            messagebox.showinfo(
                "Info",
                "Nie udało się zaktualizować kolumny kontrolnej w Excelu "
                "(sprawdź, czy plik nie jest otwarty). "
                "Używam lokalnego pliku śledzenia, więc wygenerują się "
                "tylko zmienione PDF-y.",
            )

    if excel_rows is None:
        changed_rows = cache_rows
        tracking_mode = "cache" if cache_rows is not None else tracking_mode
    else:
        changed_rows = excel_rows
        if cache_rows is not None and len(excel_rows) == total_rows and len(cache_rows) < total_rows:
            logger.info(
                "Excel tracking indicates all rows changed; using cache (%s/%s).",
                len(cache_rows),
                total_rows,
            )
            changed_rows = cache_rows
            tracking_mode = "cache"

    tracked_columns = sum(len(cols) for cols in sheet_fields_tracking.values())
    cache_count = len(cache_rows) if cache_rows is not None else None
    excel_count = len(excel_rows) if excel_rows is not None else None
    if changed_rows is None:
        logger.info(
            "Tracking mode=%s; tracked_columns=%s; cache_rows=%s; excel_rows=%s; generating all rows=%s",
            tracking_mode,
            tracked_columns,
            cache_count,
            excel_count,
            total_rows,
        )
    else:
        logger.info(
            "Tracking mode=%s; tracked_columns=%s; cache_rows=%s; excel_rows=%s; changed_rows=%s/%s",
            tracking_mode,
            tracked_columns,
            cache_count,
            excel_count,
            len(changed_rows),
            total_rows,
        )
        if len(changed_rows) >= max(1, total_rows - 1) and cache_changed_cols:
            changed_preview = []
            for sheet, cols in cache_changed_cols.items():
                for col in cols:
                    changed_preview.append(f"{sheet}:{col}")
                    if len(changed_preview) >= 6:
                        break
                if len(changed_preview) >= 6:
                    break
            logger.info(
                "Most rows changed; changed columns (sample): %s",
                ", ".join(changed_preview),
            )

    output_dir = os.path.join(os.path.dirname(app.excel_path), "PDS")
    os.makedirs(output_dir, exist_ok=True)
    report["output_dir"] = output_dir

    page_width = app.page_width
    page_height = app.page_height

    static_entries = _collect_static_entries(app)

    sheet_columns = {
        sheet: {col for col in df.columns if col != TRACKING_COLUMN}
        for sheet, df in app.dataframes.items()
    }

    name_column = _first_data_column(first_df)
    filename_counters = {}
    tasks = []
    for idx in range(total_rows):
        first_val = _get_filename_value(first_df, name_column, idx)
        filename = sanitize_filename(first_val) or f"pds_{idx + 1}"
        count = filename_counters.get(filename, 0)
        filename_counters[filename] = count + 1
        if count:
            unique_name = f"{filename}_{count + 1}"
        else:
            unique_name = filename
        pdf_path = os.path.join(output_dir, f"{unique_name}.pdf")
        existed_before = os.path.exists(pdf_path)
        if existed_before and changed_rows is not None and idx not in changed_rows:
            continue
        row_values = {}
        for sheet, columns in sheet_fields_all.items():
            df = app.dataframes.get(sheet)
            if df is None or idx >= len(df):
                for col in columns:
                    row_values[f"{sheet}:{col}"] = ""
                continue
            row_series = df.iloc[idx]
            available = sheet_columns.get(sheet, set())
            for col in columns:
                if col in available:
                    value = row_series[col]
                else:
                    value = ""
                if pd.isna(value):
                    value = ""
                else:
                    value = round_numeric_value(value)
                row_values[f"{sheet}:{col}"] = value
        tasks.append(
            {
                "idx": idx,
                "pdf_path": pdf_path,
                "name": unique_name,
                "row_values": row_values,
                "existed_before": existed_before,
                "change_details": list(cache_row_changes.get(idx) or []),
            }
        )

    if not tasks:
        try:
            mark_execution_error_rows(
                app.excel_path,
                first_sheet_name,
                [],
                total_rows=total_rows,
            )
        except Exception:
            logger.exception("Failed to clear execution error markers in Excel")
        report["total_tasks"] = 0
        report["skipped_rows"] = total_rows
        messagebox.showinfo(
            "Info", "Brak zmian - wszystkie pliki PDF są aktualne."
        )
        if report["warnings"]:
            messagebox.showwarning(
                "Uwaga",
                "Brak zmian w PDF, ale wykryto ostrzeżenia jakości danych: "
                f"{len(report['warnings'])}.",
            )
        _ui_counts(app, total_rows=total_rows, total_tasks=0, processed=0, skipped=total_rows)
        finish_now("Brak zmian", "no_changes")
        return False

    skipped_rows = total_rows - len(tasks) if changed_rows is not None else 0
    report["total_tasks"] = len(tasks)
    report["skipped_rows"] = skipped_rows
    _ui_counts(
        app,
        total_rows=total_rows,
        total_tasks=len(tasks),
        processed=0,
        skipped=skipped_rows,
    )
    _ui_progress_mode(app, False)
    _ui_status(app, "Sprawdzanie danych i generowanie PDF...")

    if _is_cancelled(app):
        finish_now("Anulowano", "cancelled")
        return False

    worker_payload = {
        "static_entries": static_entries,
        "elements": element_specs,
        "element_order": element_order,
        "groups": group_specs,
        "conditions": conditions,
        "dependencies": dependencies,
        "scale": app.scale,
        "grid_size": getattr(app, "grid_size", DEFAULT_GRID_SIZE),
        "page_width": page_width,
        "page_height": page_height,
        "excel_path": getattr(app, "excel_path", ""),
        "excel_dir": os.path.dirname(app.excel_path),
        "pdf_image_compression_percent": normalize_pdf_image_compression_percent(
            getattr(
                app,
                "pdf_image_compression_percent",
                DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
            )
        ),
        "image_fields": sorted(getattr(app, "image_fields", set())),
        "image_dirs": list(getattr(app, "image_dirs", [])),
        "tasks": tasks,
        "output_dir": output_dir,
        "total_tasks": len(tasks),
        "total_rows": total_rows,
        "skipped_rows": skipped_rows,
        "image_index_path": image_index_utils.get_default_index_path(),
    }

    def worker(payload):
        start_time = time.time()
        roots = image_index_utils.normalize_roots(
            [payload["excel_dir"]] + list(payload.get("image_dirs", []))
        )
        if roots:
            _ui_status(app, "Aktualizacja indeksu obrazów...")
            try:
                progress_state = {"last_emit": 0.0, "last_processed": 0}

                def index_progress(progress):
                    now = time.time()
                    phase = str(progress.get("phase") or "")
                    processed = int(progress.get("processed", 0) or 0)
                    if (
                        phase != "done"
                        and now - progress_state["last_emit"] < 2.0
                        and processed - int(progress_state["last_processed"]) < 1000
                    ):
                        return
                    progress_state["last_emit"] = now
                    progress_state["last_processed"] = processed
                    total_estimate = progress.get("total_estimate")
                    eta_seconds = progress.get("eta_seconds")
                    eta_text = ""
                    if eta_seconds is not None:
                        eta_text = f", ETA ~{max(0, int(eta_seconds))}s"
                    if phase == "cached":
                        status = f"Aktualizacja indeksu obrazów: użyto cache ({processed} plików)"
                    elif total_estimate:
                        status = (
                            "Aktualizacja indeksu obrazów: "
                            f"{processed}/{int(total_estimate)} plików{eta_text}"
                        )
                    else:
                        status = f"Aktualizacja indeksu obrazów: {processed} plików{eta_text}"
                    _ui_status(app, status)

                index_data = image_index_utils.ensure_index(
                    roots,
                    index_path=payload.get("image_index_path") or None,
                    max_age_hours=24,
                    force_rebuild=False,
                    progress_callback=index_progress,
                    progress_interval_seconds=1.0,
                )
                _cache_image_index_for_app(app, roots, index_data)
            except Exception:
                logger.exception("Failed to build/load image index for generation")
            _ui_status(app, "Sprawdzanie danych i generowanie PDF...")

        context = {
            "scale": payload["scale"],
            "page_width": payload["page_width"],
            "page_height": payload["page_height"],
            "excel_path": payload.get("excel_path", ""),
            "elements": payload["elements"],
            "element_order": payload["element_order"],
            "groups": payload["groups"],
            "conditions": payload["conditions"],
            "dependencies": payload.get("dependencies", []),
            "static_entries": payload["static_entries"],
            "excel_dir": payload["excel_dir"],
            "grid_size": payload.get("grid_size", DEFAULT_GRID_SIZE),
            "pdf_image_compression_percent": payload.get(
                "pdf_image_compression_percent",
                DEFAULT_PDF_IMAGE_COMPRESSION_PERCENT,
            ),
            "image_fields": payload.get("image_fields", []),
            "image_dirs": payload.get("image_dirs", []),
            "image_index_path": payload.get("image_index_path", ""),
        }

        tasks_local = payload["tasks"]
        total = payload["total_tasks"]
        failures = []
        completed = 0
        cancelled = False
        new_files = []
        updated_files = []
        updated_pdf_changes = []
        updated_pdf_changes_missing = False

        validation_warnings = []
        missing_image_issues = []
        validation_row_indices = [task.get("idx") for task in tasks_local]
        try:
            validation_warnings.extend(
                _collect_red_nonpositive_issues(
                    app,
                    validation_row_indices,
                    sheet_fields_all,
                    image_sheet_fields=image_sheet_fields,
                )
            )
        except Exception:
            logger.exception("Failed while validating warning-marked numeric fields")
            validation_warnings.append(
                "Nie udało się sprawdzić pól oznaczonych kolorem ostrzeżenia."
            )
        missing_image_rows = set()
        missing_image_issues_by_row = {}
        non_row_image_issues = []
        remote_image_limit_warning = False
        try:
            missing_image_issues = _collect_missing_image_issues(app, validation_row_indices)
        except Exception:
            logger.exception("Failed while validating image references")
            validation_warnings.append(
                "Nie udało się sprawdzić ścieżek i linków obrazów."
            )
        else:
            for issue in missing_image_issues:
                if issue == REMOTE_IMAGE_CHECK_LIMIT_WARNING:
                    remote_image_limit_warning = True
                    continue
                row_idx = _issue_row_index(issue)
                if row_idx is None:
                    non_row_image_issues.append(issue)
                else:
                    missing_image_rows.add(row_idx)
                    missing_image_issues_by_row.setdefault(row_idx, set()).add(
                        _strip_issue_row_prefix(issue)
                    )

        report["skipped_image_files"] = []
        report["skipped_image_global_issues"] = sorted(set(non_row_image_issues))
        if remote_image_limit_warning:
            validation_warnings.append(REMOTE_IMAGE_CHECK_LIMIT_WARNING)

        has_non_row_image_issue = bool(non_row_image_issues)

        skipped_due_missing_images = 0
        skipped_tasks = []
        if has_non_row_image_issue:
            skipped_due_missing_images = len(tasks_local)
            skipped_tasks = list(tasks_local)
            tasks_local = []
        elif missing_image_rows:
            filtered_tasks = []
            for task in tasks_local:
                if task.get("idx") in missing_image_rows:
                    skipped_due_missing_images += 1
                    skipped_tasks.append(task)
                else:
                    filtered_tasks.append(task)
            tasks_local = filtered_tasks

        if skipped_due_missing_images:
            total = len(tasks_local)
            report["total_tasks"] = total
            report["skipped_rows"] = (
                report.get("skipped_rows") or 0
            ) + skipped_due_missing_images
            skipped_image_files = []
            for task in skipped_tasks:
                row_no = _excel_row_number(task.get("idx") or 0)
                pdf_name = os.path.basename(task.get("pdf_path", ""))
                if task.get("existed_before"):
                    action = "pomijam aktualizację"
                else:
                    action = "pomijam utworzenie"
                row_issues = sorted(missing_image_issues_by_row.get(task.get("idx"), set()))
                if not row_issues and non_row_image_issues:
                    row_issues = sorted(set(non_row_image_issues))
                skipped_image_files.append(
                    {
                        "row": row_no,
                        "pdf_name": pdf_name,
                        "action": action,
                        "issues": row_issues,
                    }
                )

            report["skipped_image_files"] = sorted(
                skipped_image_files,
                key=lambda item: (item.get("row") or 0, item.get("pdf_name") or ""),
            )

            _ui_counts(
                app,
                total_tasks=total,
                processed=0,
                skipped=report["skipped_rows"],
            )

        if validation_warnings:
            deduped_warnings = []
            seen_warnings = set()
            for warning_text in validation_warnings:
                if warning_text in seen_warnings:
                    continue
                seen_warnings.add(warning_text)
                deduped_warnings.append(warning_text)
            report["warnings"] = deduped_warnings
            sample = ", ".join(report["warnings"][:3])
            logger.warning(
                "Detected data quality warnings: %s (sample: %s)",
                len(report["warnings"]),
                sample,
            )
            for warning_text in report["warnings"][:200]:
                logger.warning("Data warning: %s", warning_text)

        if not tasks_local:
            def finish_only_skipped():
                try:
                    mark_execution_error_rows(
                        app.excel_path,
                        first_sheet_name,
                        [],
                        total_rows=total_rows,
                    )
                except Exception:
                    logger.exception("Failed to clear execution error markers in Excel")
                report["status"] = "error"
                report["processed_rows"] = 0
                report["new_pdfs"] = []
                report["updated_pdfs"] = []
                if hasattr(app, "finish_generation_ui"):
                    app.finish_generation_ui("Błąd")
                else:
                    _ui_finish(app, "Błąd")
                messagebox.showerror(
                    "Błąd",
                    "Nie wygenerowano plików PDF, ponieważ wykryto brakujące obrazy. "
                    "Szczegóły znajdziesz w sekcji Błędy raportu i logach.",
                )
                _notify_generation_complete(app, report)

            _ui_call(app, finish_only_skipped)
            return

        max_workers = max(1, min(len(tasks_local), os.cpu_count() or 1))

        executor = None
        future_map = {}
        try:
            executor = ProcessPoolExecutor(
                max_workers=max_workers,
                initializer=_init_render_context,
                initargs=(context,),
            )
            future_map = {
                executor.submit(render_single_pdf, task): task for task in tasks_local
            }
            for future in as_completed(future_map):
                if _is_cancelled(app):
                    cancelled = True
                    break
                task = future_map[future]
                try:
                    result = future.result()
                    final_pdf = result.get("pdf_path", task["pdf_path"])
                    if task.get("existed_before"):
                        updated_files.append(final_pdf)
                        change_details = list(task.get("change_details") or [])
                        if change_details:
                            updated_pdf_changes.append(
                                {
                                    "row": _excel_row_number(task["idx"]),
                                    "pdf_name": os.path.basename(final_pdf),
                                    "pdf_path": final_pdf,
                                    "changes": change_details,
                                }
                            )
                        else:
                            updated_pdf_changes_missing = True
                    else:
                        new_files.append(final_pdf)
                except Exception as exc:  # pragma: no cover - defensive logging
                    report["critical_exception"] = True
                    failures.append((task["idx"], task.get("name", ""), str(exc)))
                    logger.exception(
                        "Failed to render PDF for row %s",
                        _excel_row_number(task["idx"]),
                        exc_info=exc,
                    )
                completed += 1
                progress = completed / total * 100
                elapsed = time.time() - start_time
                remaining = (elapsed / completed) * (total - completed) if completed else 0
                remaining_seconds = max(0, int(remaining))
                app.progress.after(
                    0, lambda p=progress: app.progress.config(value=p)
                )
                _ui_counts(
                    app,
                    total_tasks=total,
                    processed=completed,
                )
                app.time_label.after(
                    0,
                    lambda r=remaining_seconds: app.time_label.config(
                        text=f"Pozostały czas: {r} s"
                    ),
                )
        except Exception as exc:  # pragma: no cover - executor level failure
            report["critical_exception"] = True
            failures.append((-1, "", str(exc)))
            logger.exception("PDF generation failed", exc_info=exc)
        finally:
            if executor is not None:
                if cancelled:
                    for fut in future_map:
                        fut.cancel()
                    try:
                        executor.shutdown(wait=False, cancel_futures=True)
                    except TypeError:
                        try:
                            executor.shutdown(wait=False)
                        except TypeError:
                            executor.shutdown()
                else:
                    executor.shutdown()

            def finish():
                report["processed_rows"] = completed
                report["new_pdfs"] = sorted(new_files)
                report["updated_pdfs"] = sorted(updated_files)
                report["updated_pdf_changes"] = sorted(
                    updated_pdf_changes,
                    key=lambda item: (item.get("row") or 0, item.get("pdf_name") or ""),
                )
                if report["updated_pdfs"] and updated_pdf_changes_missing:
                    report["updated_pdf_changes_note"] = (
                        "Dla części zaktualizowanych PDF brak pełnej historii poprzednich "
                        "wartości, więc zestawienie zmian może być niepełne."
                    )
                for idx, name, err in failures:
                    if idx >= 0:
                        if name:
                            report["errors"].append(
                                f"Wiersz {_excel_row_number(idx)} ({name}): {err}"
                            )
                        else:
                            report["errors"].append(f"Wiersz {_excel_row_number(idx)}: {err}")
                    else:
                        report["errors"].append(f"Błąd wykonawcy: {err}")

                # Mirror execution exceptions in Excel for quick triage and clear stale markers.
                failed_rows = sorted({idx for idx, _name, _err in failures if idx >= 0})
                try:
                    mark_execution_error_rows(
                        app.excel_path,
                        first_sheet_name,
                        failed_rows,
                        total_rows=total_rows,
                    )
                except Exception:
                    logger.exception("Failed to mark execution error rows in Excel")

                if cancelled:
                    report["status"] = "cancelled"
                    if hasattr(app, "finish_generation_ui"):
                        app.finish_generation_ui("Anulowano")
                    else:
                        _ui_finish(app, "Anulowano")
                    _notify_generation_complete(app, report)
                    return
                has_missing_image_errors = bool(
                    report.get("skipped_image_files") or report.get("skipped_image_global_issues")
                )
                if failures or has_missing_image_errors:
                    report["status"] = "error"
                    if hasattr(app, "finish_generation_ui"):
                        app.finish_generation_ui("Błąd")
                    else:
                        _ui_finish(app, "Błąd")
                    if failures:
                        failed_rows = [idx for idx, _name, _err in failures if idx >= 0]
                        if failed_rows:
                            rows_text = ", ".join(
                                str(_excel_row_number(idx)) for idx in failed_rows
                            )
                            message = (
                                "Wystąpiły błędy podczas generowania wierszy: "
                                f"{rows_text}. Sprawdź logi."
                            )
                        else:
                            message = "Wystąpił błąd podczas generowania plików PDF. Sprawdź logi."
                    else:
                        skipped_count = len(report.get("skipped_image_files") or [])
                        message = (
                            "Wykryto brakujące obrazy i pominięto część zadań PDF. "
                            f"Pominięte: {skipped_count}. Sprawdź sekcję Błędy raportu i logi."
                        )
                    messagebox.showerror("Błąd", message)
                else:
                    report["status"] = "success"
                    if hasattr(app, "finish_generation_ui"):
                        app.finish_generation_ui("Zakończono")
                    else:
                        _ui_finish(app, "Zakończono")
                    messagebox.showinfo(
                        "Zakończono", f"Pliki zapisane w {payload['output_dir']}"
                    )
                    if report["warnings"]:
                        messagebox.showwarning(
                            "Uwaga",
                            "Wykryto ostrzeżenia jakości danych: "
                            f"{len(report['warnings'])}. "
                            "Szczegóły są dostępne w raporcie e-mail i logach.",
                        )
                _notify_generation_complete(app, report)

            _ui_call(app, finish)

    threading.Thread(target=worker, args=(worker_payload,), daemon=True).start()
    return True
